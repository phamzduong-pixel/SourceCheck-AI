"""DOCX parser using python-docx."""

import io
from typing import Any, Dict
import docx
from docx.opc.exceptions import PackageNotFoundError

from app.core.exceptions import DocumentParsingException
from app.services.ingestion.parsers.base import (
    BaseDocumentParser,
    PageContent,
    ParsedDocument,
)


class DocxParser(BaseDocumentParser):
    """Parser for Microsoft Word (.docx) files."""

    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        try:
            doc = docx.Document(io.BytesIO(content))
        except (PackageNotFoundError, Exception) as e:
            raise DocumentParsingException(filename, f"Invalid or corrupted DOCX file: {str(e)}")

        paragraphs_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs_text.append(p.text.strip())

        # Extract table content as well
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs_text.append(" | ".join(row_text))

        full_text = "\n\n".join(paragraphs_text)

        # Extract metadata from core properties
        docx_metadata: Dict[str, Any] = {"page_count": 1}
        try:
            core = doc.core_properties
            if core.title:
                docx_metadata["title"] = str(core.title)
            if core.author:
                docx_metadata["author"] = str(core.author)
            if core.subject:
                docx_metadata["subject"] = str(core.subject)
            if core.comments:
                docx_metadata["comments"] = str(core.comments)
        except Exception:
            pass

        pages = [PageContent(page_number=1, text=full_text)]

        return ParsedDocument(
            filename=filename,
            file_type="docx",
            pages=pages,
            full_text=full_text,
            metadata=docx_metadata,
        )
